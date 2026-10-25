#Monta o .xlsx para importar no Google Sheets, com as abas Bronze, Prata e Log
#a aba Ouro é criada no Sheets, como tabela dinâmica.

import shutil
import pandas as pd
from openpyxl import load_workbook
from src.config import BRONZE, LOG, PLANILHA_SHEETS, PRATA

def executar():
    PLANILHA_SHEETS.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(BRONZE, PLANILHA_SHEETS)  # copia o arquivo em vez de reescrever os dados: Bronze intocada
    wb = load_workbook(PLANILHA_SHEETS)
    wb.active.title = "Bronze"
    wb.save(PLANILHA_SHEETS)

    prata = pd.read_csv(PRATA, parse_dates=["Data"])
    log = pd.read_csv(LOG, dtype=str, keep_default_na=False)
    with pd.ExcelWriter(PLANILHA_SHEETS, engine="openpyxl", mode="a") as w:
        prata.to_excel(w, sheet_name="Prata", index=False)
        log.to_excel(w, sheet_name="Log", index=False)
        for (celula,) in w.sheets["Prata"].iter_rows(min_row=2, min_col=2, max_col=2):
            celula.number_format = "DD/MM/YYYY"  # mostra só a data, sem o 00:00:00
    print(f"\n[Exportar] Planilha para o Google Sheets: {PLANILHA_SHEETS.relative_to(PLANILHA_SHEETS.parents[1])}")


if __name__ == "__main__":
    executar()
